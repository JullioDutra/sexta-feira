"""Memória dos documentos: leitura, índice FTS5, busca, respostas, atalhos, pastas e rotas."""

from __future__ import annotations

import os
import sqlite3
import time
from pathlib import Path, PureWindowsPath

import pytest
from fastapi.testclient import TestClient

from sexta.cerebro.agente import Agente
from sexta.cerebro.ferramentas import Contexto
from sexta.extensoes import documentos as mod
from sexta.util.texto import normalizar

from test_agente import OllamaFalso, cliente


# ---------------------------------------------------------------- ajudantes

def criar_pdf(caminho: Path, paginas: list[str], senha: str | None = None) -> Path:
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    escritor = PdfWriter()
    fonte = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"),
                              NameObject("/BaseFont"): NameObject("/Helvetica"),
                              NameObject("/Encoding"): NameObject("/WinAnsiEncoding")})
    for texto in paginas:
        pagina = escritor.add_blank_page(612, 792)
        pagina[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): fonte})})
        conteudo = DecodedStreamObject()
        conteudo.set_data(b"BT /F1 11 Tf 72 720 Td (" + texto.encode("cp1252") + b") Tj ET")
        pagina.replace_contents(conteudo)
    if senha is not None:
        escritor.encrypt(user_password=senha, owner_password="dono")
    caminho.parent.mkdir(parents=True, exist_ok=True)
    with open(caminho, "wb") as arquivo:
        escritor.write(arquivo)
    return caminho


def criar_docx(caminho: Path, paragrafos: list[str], tabela: list[list[str]] | None = None) -> Path:
    import docx

    documento = docx.Document()
    for paragrafo in paragrafos:
        documento.add_paragraph(paragrafo)
    if tabela:
        t = documento.add_table(rows=len(tabela), cols=len(tabela[0]))
        for i, linha in enumerate(tabela):
            for j, valor in enumerate(linha):
                t.cell(i, j).text = valor
    caminho.parent.mkdir(parents=True, exist_ok=True)
    documento.save(str(caminho))
    return caminho


def escrever(caminho: Path, texto: str, codificacao: str = "utf-8") -> Path:
    caminho.parent.mkdir(parents=True, exist_ok=True)
    caminho.write_bytes(texto.encode(codificacao))
    return caminho


def sem_ia(app):
    app.agente = Agente(app, cliente(OllamaFalso([])))
    return app.agente


@pytest.fixture
def casa(tmp_path, monkeypatch):
    casa = tmp_path.resolve() / "casa"
    for nome in ("Desktop", "Documents", "Downloads", "Pictures", "Music", "Videos", "AppData"):
        (casa / nome).mkdir(parents=True)
    monkeypatch.setattr(Path, "home", classmethod(lambda cls: casa))
    monkeypatch.delenv("DOCUMENTOS_PASTAS", raising=False)
    return casa


@pytest.fixture
def memoria(app, casa):
    yield app.documentos
    app.documentos.encerrar()


@pytest.fixture
def contratos(app, casa, memoria):
    """Pasta Documentos/Contratos já na memória e indexada."""
    pasta = casa / "Documents" / "Contratos"
    criar_pdf(pasta / "contrato_aluguel.pdf", [
        "Contrato de locação residencial entre as partes.",
        "Cláusula 8: a multa rescisória é de três aluguéis, proporcional ao tempo restante.",
        "Cláusula 12: o aviso prévio para desocupação é de 30 dias.",
    ])
    escrever(pasta / "manual_geladeira.txt", "A geladeira não tem multa. Limpe o filtro a cada seis meses.")
    escrever(pasta / "iptu_2025.txt", "Prefeitura: IPTU 2025 pago em cota única de R$ 1.234,56 em fevereiro.", "cp1252")
    app.prefs.atualizar({"documentos_pastas": [str(pasta)]})
    assert memoria.indexar()
    return pasta


def respostas_falsas(app, monkeypatch, resposta: str = "A multa é de três aluguéis (contrato_aluguel.pdf, página 2)."):
    chamadas: list[tuple[str, str]] = []

    def completar(sistema, texto, rapido=True):
        chamadas.append((sistema, texto))
        return resposta

    monkeypatch.setattr(app.agente, "completar", completar)
    return chamadas


# ---------------------------------------------------------------- texto e trechos

def test_fatiar_respeita_tamanho_sobreposicao_e_paginas():
    frase = "A locatária pagará o aluguel até o quinto dia útil de cada mês, com correção anual pelo IPCA. "
    trechos = mod.fatiar(frase * 40)
    assert len(trechos) > 3
    assert all(len(t) <= mod.TAMANHO_TRECHO for t in trechos)
    assert all(t.endswith(".") for t in trechos[:-1])  # corta em fim de frase
    assert trechos[1][:30] in trechos[0]  # sobreposição entre trechos vizinhos
    assert mod.fatiar("  \n ") == [] and mod.fatiar("curto") == ["curto"]
    paginas = mod.fatiar_paginas([(1, "primeira"), (2, frase * 15), (3, "")])
    assert paginas[0] == (1, "primeira") and {p for p, _ in paginas} == {1, 2}


def test_limpar_texto_mantem_linhas():
    assert mod.limpar_texto("a\t\tb\x00c\n\n\n  d  e\x07") == "a b c\nd e"


def test_termos_da_pergunta_e_consulta_segura():
    assert mod.termos_da_pergunta("O que diz meu contrato sobre multas?") == [("contrato", True), ("multa", True)]
    assert mod.termos_da_pergunta("quanto paguei de IPTU em 2025") == [("paguei", True), ("iptu", True), ("2025", False)]
    assert mod._radical("aluguel") == mod._radical("alugueis") == "alugue"
    assert mod._radical("rescisao") == mod._radical("rescisoes") == "rescis"
    consulta = mod.consulta_fts(mod.termos_da_pergunta('multa" OR texto:* ) NEAR( "x" -iptu AND'))
    assert consulta == '"multa"* OR "or" OR "texto"* OR "near"* OR "iptu"* OR "and"'
    banco = sqlite3.connect(":memory:")
    banco.executescript(mod.ESQUEMA)
    assert banco.execute("SELECT count(*) FROM trechos WHERE trechos MATCH ?", (consulta,)).fetchone() == (0,)
    banco.close()
    assert mod.termos_da_pergunta("o que é isso?") == []


def test_destacar_vira_segmentos():
    assert mod.destacar("a \x02multa\x03 é de \x02três\x03") == [("a ", False), ("multa", True), (" é de ", False),
                                                                ("três", True)]


# ---------------------------------------------------------------- leitura dos arquivos

def test_extrair_pdf_por_pagina_e_protegido(tmp_path):
    lido = mod.extrair(criar_pdf(tmp_path / "a.pdf", ["Primeira página.", "Segunda: multa rescisória."]))
    assert lido.paginas == 2 and lido.textos == [(1, "Primeira página."), (2, "Segunda: multa rescisória.")]
    # só restrição de permissões (sem senha para abrir): lê normalmente
    assert mod.extrair(criar_pdf(tmp_path / "b.pdf", ["Aberto."], senha="")).textos == [(1, "Aberto.")]
    with pytest.raises(mod.ErroLeitura, match="senha"):
        mod.extrair(criar_pdf(tmp_path / "c.pdf", ["Segredo."], senha="1234"))
    with pytest.raises(mod.ErroLeitura, match="sem texto"):
        mod.extrair(criar_pdf(tmp_path / "d.pdf", [""]))


def test_extrair_pdf_grande_demais(tmp_path, monkeypatch):
    monkeypatch.setattr(mod, "MAX_PAGINAS", 2)
    with pytest.raises(mod.ErroLeitura, match="3 páginas"):
        mod.extrair(criar_pdf(tmp_path / "a.pdf", ["1", "2", "3"]))


def test_extrair_docx_e_texto(tmp_path):
    lido = mod.extrair(criar_docx(tmp_path / "a.docx", ["Seguro residencial", ""],
                                  [["Cobertura", "Valor"], ["Incêndio", "R$ 300.000"]]))
    assert lido.textos == [(None, "Seguro residencial\nCobertura | Valor\nIncêndio | R$ 300.000")]
    assert mod.extrair(escrever(tmp_path / "b.txt", "Aviso prévio", "cp1252")).textos == [(None, "Aviso prévio")]
    assert mod.extrair(escrever(tmp_path / "c.md", "\ufeff# Notas", "utf-8")).textos == [(None, "# Notas")]
    with pytest.raises(mod.ErroLeitura):
        mod.extrair(escrever(tmp_path / "d.csv", "  \n"))


# ---------------------------------------------------------------- índice

def test_indexar_le_pastas_e_registra_erros(app, casa, memoria, monkeypatch):
    pasta = casa / "Documents" / "Contratos"
    criar_pdf(pasta / "contrato.pdf", ["Multa de três aluguéis."])
    criar_pdf(pasta / "protegido.pdf", ["x"], senha="1234")
    escrever(pasta / "notas.md", "anotações")
    escrever(pasta / "foto.jpg", "não é documento")
    escrever(pasta / "~$contrato.docx", "temporário do Word")
    escrever(pasta / ".oculto" / "a.txt", "oculto")
    escrever(pasta / "node_modules" / "b.txt", "programa")
    escrever(pasta / "grande.txt", "x" * 200_000)
    monkeypatch.setattr(mod, "MAX_BYTES", 100_000)
    app.prefs.atualizar({"documentos_pastas": [str(pasta)]})

    assert memoria.indexar()
    estado = memoria.estado()
    assert estado["documentos"] == 2 and estado["trechos"] == 2 and estado["ultima"]
    erros = {e["arquivo"]: e["erro"] for e in estado["erros"]}
    assert erros["protegido.pdf"] == "PDF protegido por senha" and erros["grande.txt"].startswith("maior que")
    assert estado["total_erros"] == 2
    assert estado["pastas"] == [{"caminho": str(pasta), "origem": "ajustes", "erro": ""}]


def test_indexacao_incremental_e_remocao(app, contratos, memoria, monkeypatch):
    lidos = []
    original = mod.extrair
    monkeypatch.setattr(mod, "extrair", lambda caminho: lidos.append(caminho.name) or original(caminho))
    assert memoria.indexar() and lidos == []  # nada mudou: nada é relido

    escrever(contratos / "manual_geladeira.txt", "A geladeira agora tem garantia estendida de dois anos.")
    (contratos / "iptu_2025.txt").unlink()
    assert memoria.indexar() and lidos == ["manual_geladeira.txt"]
    assert memoria.estado()["documentos"] == 2
    assert [t.arquivo for t in memoria.buscar("garantia estendida")] == ["manual_geladeira.txt"]
    assert memoria.buscar("IPTU") == []

    app.prefs.atualizar({"documentos_pastas": []})  # pasta tirada da memória: os documentos dela saem do índice
    assert memoria.indexar() and memoria.estado()["documentos"] == 0


def test_pasta_indisponivel_mantem_o_indice(app, contratos, memoria):
    contratos.rename(contratos.with_name("Contratos (desconectado)"))  # ex.: disco externo desligado
    assert memoria.indexar()
    estado = memoria.estado()
    assert estado["documentos"] == 3 and estado["pastas"][0]["erro"]


def test_indexar_nao_roda_duas_vezes_ao_mesmo_tempo(memoria):
    with memoria._indexando:
        assert memoria.indexar() is False


def test_thread_de_fundo_indexa_e_encerra(app, casa, monkeypatch):
    escrever(casa / "Documents" / "nota.txt", "Senha do wifi na geladeira.")
    app.prefs.atualizar({"documentos_pastas": [str(casa / "Documents")]})
    monkeypatch.setattr(mod, "ESPERA_INICIAL_S", 0)
    mod.iniciar(app)
    try:
        limite = time.monotonic() + 10
        while app.documentos.ultima_indexacao() is None and time.monotonic() < limite:
            time.sleep(0.02)
        assert app.documentos.estado()["documentos"] == 1
    finally:
        mod.encerrar(app)
    assert not app.documentos._thread.is_alive() and app.documentos._db is None


# ---------------------------------------------------------------- busca e respostas

def test_buscar_da_bonus_ao_nome_do_arquivo(app, contratos, memoria):
    escrever(contratos / "multas_transito.txt", "Multa. Multa grave. Multa gravíssima. Multa por velocidade, multa.")
    assert memoria.indexar()
    trechos = memoria.buscar("o que diz meu contrato sobre multa?")
    assert trechos[0].arquivo == "contrato_aluguel.pdf" and trechos[0].pagina == 2
    assert trechos[0].cartao()["trecho"][1] == ("multa", True)
    assert memoria.buscar("multas de transito")[0].arquivo == "multas_transito.txt"


def test_perguntar_responde_so_com_os_trechos(app, contratos, monkeypatch):
    chamadas = respostas_falsas(app, monkeypatch)
    resultado = app.registro.executar("documentos", {"acao": "perguntar", "pergunta": "Qual a multa do contrato?"},
                                      Contexto(app, canal="voz"))
    assert resultado["ok"] and resultado["_direta"]
    assert resultado["resumo"] == "A multa é de três aluguéis (contrato_aluguel.pdf, página 2)."
    sistema, texto = chamadas[0]
    assert "SOMENTE" in sistema and "2 ou 3 frases" in sistema and "não encontrou" in sistema
    assert "Pergunta: Qual a multa do contrato?" in texto and "(contrato aluguel, página 2)" in texto
    holograma = app.hologramas.lista()[-1]
    assert holograma["tipo"] == "documentos" and not holograma["dados"]["carregando"]
    fonte = holograma["dados"]["fontes"][0]
    assert fonte["arquivo"] == "contrato_aluguel.pdf" and fonte["pagina"] == 2 and isinstance(fonte["id"], int)

    app.registro.executar("documentos", {"acao": "perguntar", "pergunta": "multa"}, Contexto(app, canal="texto"))
    assert "dois parágrafos" in chamadas[1][0] and "(contrato_aluguel.pdf, página 2)" in chamadas[1][1]


def test_perguntar_sem_resultado_nao_chama_a_ia(app, contratos, monkeypatch):
    chamadas = respostas_falsas(app, monkeypatch)
    resultado = app.registro.executar("documentos", {"acao": "perguntar", "pergunta": "qual a cor do carro?"},
                                      Contexto(app))
    assert resultado["resumo"] == "Não encontrei nada sobre isso nos seus documentos." and resultado["_direta"]
    assert chamadas == [] and app.hologramas.lista()[-1]["dados"]["fontes"] == []


def test_perguntar_sem_ia_mostra_o_trecho(app, contratos, monkeypatch):
    def falhar(*_a, **_k):
        raise RuntimeError("Ollama desligado")

    monkeypatch.setattr(app.agente, "completar", falhar)
    resultado = app.registro.executar("documentos", {"acao": "perguntar", "pergunta": "aviso prévio"},
                                      Contexto(app, canal="texto"))
    assert resultado["resumo"] == ("Não consegui falar com a IA agora. O trecho mais próximo está em "
                                   "contrato_aluguel.pdf, página 3: Cláusula 12: o aviso prévio para desocupação é de 30 dias.")


def test_perguntar_sem_pasta_explica_como_configurar(app, casa, memoria):
    resultado = app.registro.executar("documentos", {"acao": "perguntar", "pergunta": "multa"}, Contexto(app))
    assert not resultado["ok"] and resultado["_direta"] and "usa a pasta" in resultado["resumo"]


def test_perguntar_antes_da_primeira_leitura(app, casa, memoria, monkeypatch):
    pedidos = []
    monkeypatch.setattr(memoria, "pedir_indexacao", lambda: pedidos.append(1))
    app.prefs.atualizar({"documentos_pastas": [str(casa / "Documents")]})
    resultado = app.registro.executar("documentos", {"acao": "perguntar", "pergunta": "multa"}, Contexto(app))
    assert "Ainda não li" in resultado["resumo"] and pedidos == [1]


def test_status_e_indexar_pela_ferramenta(app, contratos, memoria, monkeypatch):
    status = app.registro.executar("documentos", {"acao": "status"}, Contexto(app))["resumo"]
    assert status.startswith("Tenho 3 documentos") and "1 pasta na memória" in status
    pedidos = []
    monkeypatch.setattr(memoria, "pedir_indexacao", lambda: pedidos.append(1))
    assert app.registro.executar("documentos", {"acao": "indexar"}, Contexto(app))["ok"] and pedidos == [1]


# ---------------------------------------------------------------- pastas

def test_resolver_pasta_por_nome(casa):
    alvo = casa / "Documents" / "Contratos Públicos"
    alvo.mkdir()
    for texto in ("Documentos/Contratos Públicos", "documentos barra contratos publicos",
                  "contratos publicos dentro de documentos", "Contratos Públicos", str(alvo)):
        assert mod.resolver_pasta(texto) == alvo, texto
    with pytest.raises(mod.ErroDocumentos, match="Não encontrei a pasta"):
        mod.resolver_pasta("Documentos/Inexistente")


@pytest.mark.parametrize("caminho,motivo", [
    (r"C:\Users\Ana\Documents\Contratos", None),
    (r"c:\users\ana\Documents", None),
    (r"D:\Arquivos\Contratos", None),
    (r"\\nas\familia\documentos", None),
    (r"C:\Users\Ana", "não ela inteira"),
    (r"C:\Users\Ana\AppData\Roaming", "AppData"),
    (r"C:\Windows\System32", "pasta de usuário ou de outros discos"),
    (r"C:\Users\Beto\Documents", "pasta de usuário ou de outros discos"),
    ("D:\\", "não o disco inteiro"),
    (r"D:\Program Files\App", "pastas do sistema"),
    (r"E:\$RECYCLE.BIN\x", "pastas do sistema"),
])
def test_motivo_recusa_no_windows(caminho, motivo):
    resultado = mod.motivo_recusa(PureWindowsPath(caminho), PureWindowsPath(r"C:\Users\Ana"), "C:")
    assert resultado is None if motivo is None else motivo in resultado


def test_pastas_por_voz(app, casa, memoria):
    (casa / "Documents" / "Contratos").mkdir()
    agente = sem_ia(app)
    assert agente.processar("usa a pasta Documentos/Contratos como memória")["caminho"] == "atalho"
    assert app.prefs.get("documentos_pastas") == [str(casa / "Documents" / "Contratos")]
    resposta = agente.processar("adiciona a pasta de contratos na memória dos documentos")["texto"]
    assert "já está na memória" in resposta
    assert "Por segurança" in agente.processar("adiciona a pasta AppData na memória dos documentos")["texto"]
    assert agente.processar("tira a pasta contratos da memória")["texto"] == "Tirei a pasta Contratos da memória dos documentos."
    assert app.prefs.get("documentos_pastas") == []


def test_validar_pasta_recusa_fora_de_casa(casa, tmp_path):
    fora = tmp_path / "fora"
    fora.mkdir()
    with pytest.raises(mod.ErroDocumentos, match="Por segurança"):
        mod.validar_pasta(str(fora))
    with pytest.raises(mod.ErroDocumentos, match="não ela inteira"):
        mod.validar_pasta(str(casa))
    with pytest.raises(mod.ErroDocumentos, match="Não encontrei"):
        mod.validar_pasta(str(casa / "Documents" / "nada"))


def test_pasta_mae_substitui_as_de_dentro(app, casa, memoria):
    (casa / "Documents" / "Contratos").mkdir()
    mod.adicionar_pasta(app, "Documentos/Contratos")
    assert "Pronto" in mod.adicionar_pasta(app, "Documentos")
    assert app.prefs.get("documentos_pastas") == [str(casa / "Documents")]
    assert "dentro de Documents" in mod.adicionar_pasta(app, "Documentos/Contratos")


def test_preferencia_malformada_nao_quebra(app, casa, memoria):
    app.prefs.atualizar({"documentos_pastas": 5})
    assert memoria.pastas_configuradas() == []
    app.prefs.atualizar({"documentos_pastas": ["", 3, str(casa / "Documents")]})
    assert [p["caminho"] for p in memoria.pastas_configuradas()] == [str(casa / "Documents")]


def test_pastas_do_env_entram_mas_nao_saem_por_voz(app, casa, memoria, monkeypatch):
    monkeypatch.setenv("DOCUMENTOS_PASTAS", f"Documentos, {casa / 'Desktop'}")
    assert [p["origem"] for p in memoria.pastas_configuradas()] == ["env", "env"]
    raizes, _ = memoria.pastas_validas()
    assert raizes == [casa / "Documents", casa / "Desktop"]
    with pytest.raises(mod.ErroDocumentos, match=".env"):
        mod.remover_pasta(app, "Desktop")


# ---------------------------------------------------------------- atalhos

@pytest.mark.parametrize("frase", [
    "O que diz meu contrato sobre multa?",
    "Sexta-feira, o que fala a apólice sobre roubo?",
    "o que consta no meu contrato sobre reajuste",
    "o que meu contrato diz sobre multa",
    "No contrato de aluguel, qual o prazo de aviso prévio?",
    "procura nos meus documentos quanto paguei de IPTU",
    "nos meus documentos tem algo sobre seguro?",
    "pergunta pros meus documentos qual o valor do condomínio",
    "segundo meus documentos quando vence a garantia",
])
def test_atalhos_de_pergunta(app, casa, memoria, monkeypatch, frase):
    chamadas = []
    monkeypatch.setattr(app.registro, "executar",
                        lambda nome, a, ctx, **_: chamadas.append((nome, a)) or {"ok": True, "resumo": "ok"})
    assert sem_ia(app).processar(frase)["caminho"] == "atalho"
    nome, args = chamadas[0]
    assert nome == "documentos" and args["acao"] == "perguntar"
    assert args["pergunta"] == frase.removeprefix("Sexta-feira, ")


@pytest.mark.parametrize("frase,args", [
    ("reindexa os meus documentos", {"acao": "indexar"}),
    ("atualiza a memória dos documentos", {"acao": "indexar"}),
    ("usa a pasta Documentos barra Contratos como memória", {"acao": "adicionar_pasta", "pasta": "documentos barra contratos"}),
    ("remove a pasta contratos da memória dos documentos", {"acao": "remover_pasta", "pasta": "contratos"}),
])
def test_atalhos_de_pastas(app, casa, memoria, frase, args):
    rodar = lambda nome, a, padrao=None: (nome, a)
    assert mod.atalho(normalizar(frase), Contexto(app), rodar, frase) == ("documentos", args)


def test_atalho_abre_o_holograma(app, casa, memoria):
    rodar = lambda nome, a, padrao=None: (nome, a)
    assert mod.atalho("abre a memoria dos documentos", Contexto(app), rodar, None) == (
        "holograma", {"acao": "mostrar", "tipo": "documentos"})


def test_atalho_pelo_nome_de_um_documento(app, contratos):
    ctx = Contexto(app)
    rodar = lambda nome, args, padrao=None: (nome, args)
    assert mod.atalho("o que diz meu iptu sobre o valor", ctx, rodar, None) == (
        "documentos", {"acao": "perguntar", "pergunta": "o que diz meu iptu sobre o valor"})
    assert mod.atalho("o que diz o jornal sobre o cruzeiro", ctx, rodar, None) is None


@pytest.mark.parametrize("frase", [
    "abre o spotify", "fecha o chrome", "toca legião urbana", "modo jogo", "volume 30", "analisa a tela",
    "o que é isso?", "planeja meu dia", "abre o jornal", "abre o pdf do contrato de março", "arquivos recentes",
    "mostra os meus documentos recentes", "pesquisa receita de pão de queijo", "o que diz o jornal sobre o Cruzeiro",
    "o que fala a bíblia sobre perdão", "acha o contrato de aluguel", "anota: renovar o seguro", "que horas são",
    "me avisa quando sair notícia de RTX 60", "coloca o vs code na esquerda", "o que está pesando no pc?",
])
def test_atalho_nao_engole_comandos_do_nucleo(app, contratos, frase):
    t = normalizar(frase)
    assert mod.atalho(t, Contexto(app), lambda *a, **k: a, frase) is None, frase


# ---------------------------------------------------------------- HUD, rotas e diagnóstico

def test_dados_do_holograma_e_diagnostico(app, contratos, monkeypatch):
    respostas_falsas(app, monkeypatch)
    app.hologramas.mostrar("documentos")
    dados = app.hologramas.lista()[-1]["dados"]
    assert dados["estado"]["documentos"] == 3 and dados["estado"]["pastas"] == 1 and "resposta" not in dados
    app.registro.executar("documentos", {"acao": "perguntar", "pergunta": "multa"}, Contexto(app))
    assert mod.dados_holograma(app, "documentos")["resposta"].startswith("A multa")
    assert mod.dados_holograma(app, "clima") is None
    item = mod.diagnostico(app)[0]
    assert item["nome"] == "Memória dos documentos" and item["ok"] and "3 documentos" in item["detalhe"]


def test_diagnostico_sem_pastas(app, casa, memoria):
    item = mod.diagnostico(app)[0]
    assert item["ok"] and item["detalhe"] == "nenhuma pasta escolhida"


@pytest.fixture
def api(app):
    from sexta.servidor import criar_api

    with TestClient(criar_api(app), base_url="http://127.0.0.1:8765") as c:
        c.headers["Authorization"] = f"Bearer {app.acesso.token_pc}"
        yield c


def test_rotas_de_ajustes(app, casa, memoria, api):
    (casa / "Documents" / "Contratos").mkdir()
    assert api.get("/api/documentos/estado", headers={"Authorization": ""}).status_code == 401
    r = api.post("/api/documentos/pastas", json={"acao": "adicionar", "pasta": "Documentos/Contratos"})
    assert r.status_code == 200 and r.json()["estado"]["pastas"][0]["caminho"] == str(casa / "Documents" / "Contratos")
    r = api.post("/api/documentos/pastas", json={"acao": "adicionar", "pasta": "/etc"})
    assert r.status_code == 400 and "Por segurança" in r.json()["erro"]
    assert api.post("/api/documentos/pastas", json={"acao": "apagar", "pasta": "x"}).status_code == 400
    assert api.get("/api/documentos/estado").json()["documentos"] == 0
    r = api.post("/api/documentos/pastas", json={"acao": "remover", "pasta": "Contratos"})
    assert r.status_code == 200 and r.json()["estado"]["pastas"] == []


def test_rota_abrir_so_documentos_indexados(app, contratos, memoria, api, monkeypatch):
    assert api.post("/api/documentos/abrir", json={"id": "x"}).status_code == 400
    assert api.post("/api/documentos/abrir", json={"id": 999}).status_code == 404
    doc = next(t.doc for t in memoria.buscar("multa") if t.arquivo == "contrato_aluguel.pdf")
    assert api.post("/api/documentos/abrir", json={"id": doc}).status_code == 501  # fora do Windows
    abertos = []
    monkeypatch.setattr(mod.windows, "_exigir_windows", lambda: None)
    monkeypatch.setattr(os, "startfile", abertos.append, raising=False)
    assert api.post("/api/documentos/abrir", json={"id": doc}).json()["texto"] == "Abrindo contrato_aluguel.pdf."
    assert abertos == [str(contratos / "contrato_aluguel.pdf")]
    monkeypatch.setattr(app.sessao, "bloqueada", True)
    assert api.post("/api/documentos/abrir", json={"id": doc}).status_code == 423


def test_rota_perguntar(app, contratos, api, monkeypatch):
    respostas_falsas(app, monkeypatch)
    r = api.post("/api/documentos/perguntar", json={"pergunta": "qual a multa?"})
    assert r.json() == {"texto": "A multa é de três aluguéis (contrato_aluguel.pdf, página 2).", "ok": True}
    assert api.post("/api/documentos/perguntar", json={"pergunta": " "}).status_code == 400
